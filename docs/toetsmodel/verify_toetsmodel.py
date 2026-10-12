#!/usr/bin/env python3
"""Vergelijkt de (herberekende) Excel-uitkomsten met de app-engine.

Gebruik:  python3 verify_toetsmodel.py [xlsx] [layout.json] [expected.json] [scenarios.json]
Zonder argumenten: het meegeleverde werkboek met de vaste testscenario's.
"""
import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
a = sys.argv[1:]
xlsx = Path(a[0]) if len(a) > 0 else HERE / "Hypotheek_Toetsmodel_2026.xlsx"
layout = Path(a[1]) if len(a) > 1 else HERE / "layout.json"
expected = Path(a[2]) if len(a) > 2 else HERE / "expected.json"
scenarios = Path(a[3]) if len(a) > 3 else HERE / "scenarios.json"

wb = load_workbook(xlsx, data_only=True)
lay = json.loads(layout.read_text())
exp = json.loads(expected.read_text())
scn = json.loads(scenarios.read_text())
ws = wb["Berekening"]
bad = []
for i, e in enumerate(exp):
    col = lay["COLS"][i]
    for key, bkey in lay["METRICS"]:
        want = e.get(key)
        got = ws[f"{col}{lay['BR'][bkey]}"].value
        if want is None:
            ok = got == "n.v.t."
        else:
            ok = isinstance(got, (int, float)) and abs(got - want) <= (1e-6 if abs(want) < 10 else 0.01)
        if not ok:
            bad.append((i, key, got, want))
print(f"scenario's: {len(exp)}  vergelijkingen: {len(exp) * len(lay['METRICS'])}  afwijkingen: {len(bad)}")
if bad:
    print(Counter(k for _, k, _, _ in bad))
    for i, key, got, want in bad[:25]:
        print(f"  scenario {i+1} ({scn[i]['name']}): {key}: excel={got} app={want}")
sys.exit(1 if bad else 0)
